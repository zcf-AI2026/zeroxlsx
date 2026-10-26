import datetime as dt
import io
import zipfile
import zlib

import openpyxl
import pytest

from zeroxlsx import Style, Workbook
from zeroxlsx._images import probe


def make_png(width: int, height: int) -> bytes:
    """A real (if boring) RGB PNG, built without Pillow."""

    def chunk(kind: bytes, data: bytes) -> bytes:
        return (len(data).to_bytes(4, "big") + kind + data
                + (zlib.crc32(kind + data) & 0xFFFFFFFF).to_bytes(4, "big"))

    header = width.to_bytes(4, "big") + height.to_bytes(4, "big") + bytes([8, 2, 0, 0, 0])
    raw = b"".join(b"\x00" + b"\x11\x22\x33" * width for _ in range(height))
    return (b"\x89PNG\r\n\x1a\n" + chunk(b"IHDR", header)
            + chunk(b"IDAT", zlib.compress(raw)) + chunk(b"IEND", b""))


def load(data: bytes):
    return openpyxl.load_workbook(io.BytesIO(data))


# --- values ----------------------------------------------------------------

def test_writes_plain_values():
    book = Workbook()
    sheet = book.add_sheet("Data")
    sheet.append(["name", "qty", "price", "ok"])
    sheet.append(["widget", 12, 19.5, True])
    loaded = load(book.to_bytes())
    rows = list(loaded["Data"].iter_rows(values_only=True))
    assert rows[0] == ("name", "qty", "price", "ok")
    assert rows[1] == ("widget", 12, 19.5, True)


def test_write_accepts_column_letters():
    book = Workbook()
    sheet = book.add_sheet()
    sheet.write(3, "C", "cell C3")
    ws = load(book.to_bytes())["Sheet1"]
    assert ws["C3"].value == "cell C3"


def test_formula_is_written_as_formula():
    book = Workbook()
    sheet = book.add_sheet()
    sheet.append([1, 2])
    sheet.write(1, 3, "=SUM(A1:B1)")
    ws = load(book.to_bytes())["Sheet1"]
    assert ws["C1"].value == "=SUM(A1:B1)"


def test_dates_become_excel_serials():
    book = Workbook()
    sheet = book.add_sheet()
    sheet.append([dt.date(2026, 3, 1)])
    ws = load(book.to_bytes())["Sheet1"]
    cell = ws["A1"]
    assert cell.value.date() == dt.date(2026, 3, 1)
    assert cell.number_format == "yyyy-mm-dd"


def test_text_needing_escaping_survives():
    book = Workbook()
    sheet = book.add_sheet()
    tricky = 'a & b < c > d "quoted" 中文'
    sheet.append([tricky, "  padded  ", None])
    ws = load(book.to_bytes())["Sheet1"]
    assert ws["A1"].value == tricky
    assert ws["B1"].value == "  padded  "
    assert ws["C1"].value is None


# --- styles ----------------------------------------------------------------

def test_bold_and_fill_and_number_format():
    book = Workbook()
    sheet = book.add_sheet()
    header = Style(bold=True, fill="#FFCC00", h_align="center")
    sheet.append(["total"], style=header)
    sheet.append([1234.5], style=Style(number_format="#,##0.00"))
    ws = load(book.to_bytes())["Sheet1"]
    assert ws["A1"].font.bold is True
    assert ws["A1"].fill.fgColor.rgb.endswith("FFCC00")
    assert ws["A1"].alignment.horizontal == "center"
    assert ws["A2"].number_format == "#,##0.00"


def test_identical_styles_share_one_xf():
    book = Workbook()
    sheet = book.add_sheet()
    style = Style(bold=True)
    for row in range(1, 6):
        sheet.append([f"r{row}"], style=style)
    book.to_bytes()
    assert len(book.styles._xfs) == 2  # default + the one bold format


def test_style_merged_helper():
    base = Style(bold=True)
    italic = base.merged(italic=True)
    assert italic.bold and italic.italic


# --- layout ----------------------------------------------------------------

def test_column_widths_and_merges():
    book = Workbook()
    sheet = book.add_sheet()
    sheet.set_column_width("A", 24)
    sheet.set_column_width(2, 12.5)
    sheet.merge("A1:C1", "report title", Style(bold=True))
    ws = load(book.to_bytes())["Sheet1"]
    assert ws.column_dimensions["A"].width == pytest.approx(24)
    assert ws.column_dimensions["B"].width == pytest.approx(12.5)
    assert "A1:C1" in [str(rng) for rng in ws.merged_cells.ranges]
    assert ws["A1"].value == "report title"


# --- sheets ----------------------------------------------------------------

def test_multiple_sheets_and_auto_names():
    book = Workbook()
    book.add_sheet("First")
    book.add_sheet("Second")
    book.add_sheet()
    loaded = load(book.to_bytes())
    assert loaded.sheetnames == ["First", "Second", "Sheet1"]


def test_sheet_name_validation():
    book = Workbook()
    book.add_sheet("ok")
    with pytest.raises(ValueError):
        book.add_sheet("ok")          # duplicate, case-insensitive
    with pytest.raises(ValueError):
        book.add_sheet("bad/name")    # forbidden character
    with pytest.raises(ValueError):
        book.add_sheet("x" * 32)      # too long


# --- images ----------------------------------------------------------------

def test_embedded_png_round_trips():
    book = Workbook()
    sheet = book.add_sheet("Pics")
    sheet.append(["logo below"])
    sheet.add_image(make_png(40, 20), anchor="B2")
    data = book.to_bytes()

    with zipfile.ZipFile(io.BytesIO(data)) as archive:
        names = archive.namelist()
    assert "xl/media/image1.png" in names
    assert "xl/drawings/drawing1.xml" in names
    assert "xl/drawings/_rels/drawing1.xml.rels" in names

    ws = openpyxl.load_workbook(io.BytesIO(data))["Pics"]
    assert len(ws._images) == 1
    # openpyxl reports pixel size on the image; the EMU geometry lives on the anchor.
    # 40x20 px at 96 dpi -> 381000 x 190500 EMU
    anchor = ws._images[0].anchor
    assert round(anchor.ext.cx) == 381000
    assert round(anchor.ext.cy) == 190500


def test_image_size_can_be_overridden():
    book = Workbook()
    sheet = book.add_sheet()
    sheet.add_image(make_png(40, 20), anchor="A1", width=80)
    ws = openpyxl.load_workbook(io.BytesIO(book.to_bytes()))["Sheet1"]
    anchor = ws._images[0].anchor
    assert round(anchor.ext.cx) == 762000   # 80 px
    assert round(anchor.ext.cy) == 381000   # aspect ratio kept (40 px)


def test_jpeg_header_is_parsed():
    # SOI + SOF0(0xFFC0) with 300x200 + EOI — enough for the dimension probe.
    sof = b"\xff\xc0" + (17).to_bytes(2, "big") + bytes([8]) + \
        (200).to_bytes(2, "big") + (300).to_bytes(2, "big") + b"\x03" + b"\x00" * 9
    data = b"\xff\xd8" + sof + b"\xff\xd9"
    extension, width, height = probe(data)
    assert (extension, width, height) == ("jpeg", 300, 200)


def test_unsupported_image_is_rejected():
    book = Workbook()
    sheet = book.add_sheet()
    with pytest.raises(ValueError):
        sheet.add_image(b"GIF89a-not-supported")


# --- output ----------------------------------------------------------------

def test_save_to_path_and_stream(tmp_path):
    book = Workbook()
    book.add_sheet().append(["x"])
    target = tmp_path / "out.xlsx"
    book.save(target)
    assert target.stat().st_size > 1000
    assert load(target.read_bytes())["Sheet1"]["A1"].value == "x"

    stream = io.BytesIO()
    book.save(stream)
    assert stream.getvalue()[:2] == b"PK"


def test_package_contains_required_parts():
    book = Workbook()
    book.add_sheet().append(["x"])
    with zipfile.ZipFile(io.BytesIO(book.to_bytes())) as archive:
        names = set(archive.namelist())
    for required in ("[Content_Types].xml", "_rels/.rels", "xl/workbook.xml",
                     "xl/_rels/workbook.xml.rels", "xl/styles.xml",
                     "xl/worksheets/sheet1.xml"):
        assert required in names
